import os
import io
from typing import Tuple, List, Dict, Any

def extract_text_from_file(file_path: str) -> Tuple[str, List[Dict[str, Any]]]:
    """
    Extracts plain text and page/section metadata from supported file types.
    Supported types: .pdf, .docx, .txt, .md
    
    Returns:
    - full_text: string containing all extracted text
    - sections: list of dicts with text and page/section metadata
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()

    if ext in [".txt", ".md"]:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        return content, [{"section": "body", "page": 1, "text": content}]

    elif ext == ".pdf":
        return _extract_from_pdf(file_path)

    elif ext == ".docx":
        return _extract_from_docx(file_path)

    else:
        raise ValueError(f"Unsupported file format for text extraction: {ext}")

def _extract_from_pdf(file_path: str) -> Tuple[str, List[Dict[str, Any]]]:
    try:
        import pypdf
        reader = pypdf.PdfReader(file_path)
        sections = []
        full_text_parts = []

        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                sections.append({
                    "page": page_idx + 1,
                    "text": page_text
                })
                full_text_parts.append(page_text)

        full_text = "\n\n".join(full_text_parts)
        return full_text, sections
    except Exception as e:
        raise RuntimeError(f"Error parsing PDF file: {str(e)}")

def _extract_from_docx(file_path: str) -> Tuple[str, List[Dict[str, Any]]]:
    try:
        import docx
        doc = docx.Document(file_path)
        paragraphs = []
        sections = []

        for p_idx, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if text:
                paragraphs.append(text)
                sections.append({
                    "paragraph_index": p_idx,
                    "style": para.style.name if para.style else "Normal",
                    "text": text
                })

        full_text = "\n\n".join(paragraphs)
        return full_text, sections
    except Exception as e:
        raise RuntimeError(f"Error parsing DOCX file: {str(e)}")

def create_docx_file(content: str, title: str = "Document") -> bytes:
    """
    Creates a downloadable .docx file in memory from markdown/plain text content.
    """
    import docx
    from docx.shared import Pt, Inches, RGBColor

    doc = docx.Document()

    # Set document title
    title_p = doc.add_heading(title, level=0)
    
    # Process lines
    lines = content.split("\n")
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            doc.add_paragraph(stripped[2:], style='List Bullet')
        else:
            doc.add_paragraph(stripped)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()
